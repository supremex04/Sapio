import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches

from tests.conftest import png_bytes
from ved_inference.config import ParsingConfig
from ved_inference.parsing.parser import DocumentParser


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Cell Biology", level=1)
    doc.add_paragraph("Cells are the basic unit of life. ").add_run("Osmosis").bold = True
    doc.add_heading("1.1 Membranes", level=2)
    doc.add_paragraph("The membrane controls what enters and leaves the cell.")
    doc.add_paragraph("Lipid bilayer", style="List Bullet")
    doc.add_paragraph("Protein channels", style="List Bullet")
    doc.add_picture(io.BytesIO(png_bytes()), width=Inches(3))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("2 Energy", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Process", "Where"
    table.cell(1, 0).text, table.cell(1, 1).text = "Photosynthesis", "Chloroplast"
    doc.add_paragraph("Plants convert light into chemical energy in the chloroplast.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_word_documents_keep_their_own_structure():
    parsed = DocumentParser(None, ParsingConfig()).parse(make_docx(), "biology.docx")
    assert parsed.engine == "docx" and parsed.page_count == 2 and parsed.ocr_pages == 0
    titles = [(n.path, n.title) for n in parsed.tree.nodes]
    assert titles == [("1", "Cell Biology"), ("1.1", "Membranes"), ("2", "Energy")]
    text = " ".join(c.text for c in parsed.tree.chunks)
    assert "**Osmosis**" in text and "- Lipid bilayer" in text
    assert "|Process|Where|" in text and "|Photosynthesis|Chloroplast|" in text
    assert len(parsed.tree.figures) == 1 and parsed.tree.figures[0].page == 1
    energy = next(n for n in parsed.tree.nodes if n.title == "Energy")
    assert energy.page_from == 2


def test_a_corrupt_word_file_is_rejected():
    import pytest

    from ved_inference.errors import DocumentError

    with pytest.raises(DocumentError, match="Word document"):
        DocumentParser(None, ParsingConfig()).parse(b"PK\x03\x04garbage word/", "notes.docx")
